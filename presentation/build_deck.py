"""
Build the RCCG LP63 YAYA "Website & App - How To Use It" PowerPoint deck.

Run:  C:/Python313/python.exe presentation/build_deck.py
Out:  presentation/RCCG-LP63-YAYA-Website-and-App-Guide.pptx
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
ASSETS = HERE / "assets"
OUT = HERE / "RCCG-LP63-YAYA-Website-and-App-Guide.pptx"

# ---------------------------------------------------------------- theme ----
BG = RGBColor(0x0E, 0x0D, 0x18)
PANEL = RGBColor(0x18, 0x16, 0x28)
PANEL2 = RGBColor(0x21, 0x1E, 0x36)
GOLD = RGBColor(0xE8, 0xB9, 0x5F)
GOLD_DEEP = RGBColor(0xB8, 0x8A, 0x33)
GOLD_SOFT = RGBColor(0xF3, 0xD9, 0xA0)
INK = RGBColor(0xF5, 0xF2, 0xEA)
DIM = RGBColor(0xAB, 0xA6, 0xBE)
LINE = RGBColor(0x35, 0x31, 0x51)

FONT = "Segoe UI"
SW, SH = 13.333, 7.5
M = 0.75  # page margin

WEBSITE = "rccglp63yaya.org.ng"
APP_URL = "https://play.google.com/store/apps/details?id=com.churchmedia.app"
PHONES = "+234 817 812 4518 · +234 806 067 5824 · +234 811 987 1010"

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------- helpers ----
def slide(page: bool = True):
    s = prs.slides.add_slide(BLANK)
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = BG
    return s


def rect(s, x, y, w, h, fill=PANEL, line=None, lw=1.0, rounded=True, adj=0.06):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = s.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = adj
        except (IndexError, ValueError):
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.text_frame.word_wrap = True
    return shp


def tframe(container, pad=0.14):
    tf = container.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(pad * 0.8)
    return tf


def txt(s, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def par(tf, content, size=16, color=INK, bold=False, italic=False, font=FONT,
        align=PP_ALIGN.LEFT, before=0, after=0, line=None, bullet=None,
        bullet_color=GOLD, bullet_size=None, spacing=None):
    """content: str, or list of (text, bold) / (text, bold, color) tuples."""
    p = tf.paragraphs[0] if (len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    if before:
        p.space_before = Pt(before)
    if after:
        p.space_after = Pt(after)
    if line:
        p.line_spacing = line

    if bullet:
        r = p.add_run()
        r.text = bullet + "   "
        r.font.size = Pt(bullet_size or size)
        r.font.bold = True
        r.font.color.rgb = bullet_color
        r.font.name = font

    parts = [(content, False)] if isinstance(content, str) else content
    for item in parts:
        if len(item) == 3:
            t, b, c = item
        elif len(item) == 2:
            (t, b), c = item, color
        else:
            t, b, c = item[0], bold, color
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = b
        r.font.italic = italic
        r.font.color.rgb = c
        r.font.name = font
        if spacing:
            r.font._rPr.set("spc", str(int(spacing * 100)))
    return p


def header(s, eyebrow, title, sub=None, title_size=30):
    rect(s, M, 0.58, 0.055, 0.34, fill=GOLD, rounded=False)
    tf = txt(s, M + 0.18, 0.58, 10.5, 0.34, anchor=MSO_ANCHOR.MIDDLE)
    par(tf, eyebrow.upper(), size=11.5, color=GOLD, bold=True, spacing=1.6)
    tf = txt(s, M, 0.96, SW - 2 * M, 0.72)
    par(tf, title, size=title_size, color=INK, bold=True, line=0.95)
    rect(s, M, 1.74, 0.95, 0.045, fill=GOLD_DEEP, rounded=False)
    if sub:
        tf = txt(s, M, 1.9, SW - 2 * M - 0.5, 0.45)
        par(tf, sub, size=13, color=DIM, line=1.25)


def footer(s, num):
    rect(s, M, 6.9, SW - 2 * M, 0.012, fill=LINE, rounded=False)
    tf = txt(s, M, 7.02, 8.0, 0.3, anchor=MSO_ANCHOR.MIDDLE)
    par(tf, "RCCG LP63 YAYA  ·  Website & Mobile App Guide", size=9.5, color=RGBColor(0x6E, 0x69, 0x88))
    tf = txt(s, SW - M - 1.0, 7.02, 1.0, 0.3, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    par(tf, f"{num:02d}", size=10, color=GOLD_DEEP, bold=True)


def picture(s, name, x, y, w, border=True):
    path = ASSETS / name
    with Image.open(path) as im:
        iw, ih = im.size
    h = w * ih / iw
    if border:
        pad = 0.045
        rect(s, x - pad, y - pad, w + 2 * pad, h + 2 * pad, fill=None, line=LINE, lw=1.0, adj=0.02)
    s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return h


def card(s, x, y, w, h, icon, title, lines, icon_color=GOLD, title_size=15, body_size=12.5):
    rect(s, x, y, w, h, fill=PANEL, line=LINE, lw=1.0)
    if icon:
        ic = rect(s, x + 0.28, y + 0.26, 0.52, 0.52, fill=RGBColor(0x2A, 0x25, 0x42), line=RGBColor(0x4A, 0x41, 0x66), adj=0.22)
        tf = tframe(ic)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        par(tf, icon, size=19, color=icon_color, align=PP_ALIGN.CENTER)
        ty = y + 0.95
    else:
        ty = y + 0.3
    tf = txt(s, x + 0.28, ty, w - 0.56, 0.35)
    par(tf, title, size=title_size, color=INK, bold=True)
    tf = txt(s, x + 0.28, ty + 0.38, w - 0.56, h - (ty - y) - 0.5)
    for ln in lines:
        par(tf, ln, size=body_size, color=DIM, line=1.22, after=5, bullet="•",
            bullet_color=GOLD, bullet_size=body_size * 0.8)


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text.strip()


def set_table_style(tbl, style_id="{2D5ABB26-0587-4C30-8999-92F81FD0307C}"):
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    tblPr.append(parse_xml(f'<a:tableStyleId {nsdecls("a")}>{style_id}</a:tableStyleId>'))


def table(s, x, y, w, h, data, widths, row_h=0.62):
    rows, cols = len(data), len(data[0])
    gf = s.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gf.table
    set_table_style(tbl)
    total = sum(widths)
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Emu(int(Inches(w) * cw / total))
    for r in range(rows):
        tbl.rows[r].height = Inches(row_h if r else row_h * 0.9)
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = ""
            cell.margin_left = cell.margin_right = Inches(0.14)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = RGBColor(0x2A, 0x25, 0x42)
            else:
                cell.fill.fore_color.rgb = PANEL if r % 2 else PANEL2
            tf = cell.text_frame
            tf.word_wrap = True
            is_head = r == 0
            is_task = c == 0
            par(tf, data[r][c],
                size=12 if is_head or not is_task else 12.5,
                color=GOLD if is_head else (INK if is_task else DIM),
                bold=is_head or is_task, line=1.1)
    return tbl


def site_label(s, x, y):
    tf = txt(s, x, y, 6.0, 0.3)
    par(tf, WEBSITE, size=12, color=GOLD_SOFT, bold=True)
    return tf


def app_label(s, x, y, w=6.0):
    tf = txt(s, x, y, w, 0.3)
    par(tf, "CHURCH MEDIA  ·  Google Play", size=12, color=GOLD_SOFT, bold=True)


# =============================================================== SLIDES ====
# 1 ------------------------------------------------------------ title -----
s = slide(page=False)
rect(s, 0, 0, SW, 0.14, fill=GOLD_DEEP, rounded=False)

# decorative blocks
rect(s, 9.4, 0.0, 3.95, 7.5, fill=RGBColor(0x14, 0x12, 0x24), rounded=False)
rect(s, 9.4, 0.0, 0.04, 7.5, fill=RGBColor(0x2A, 0x25, 0x42), rounded=False)

tf = txt(s, 1.0, 1.55, 8.0, 0.4)
par(tf, "RCCG LP63 YAYA  ·  YAYA PROVINCE HEADQUARTERS", size=12.5, color=GOLD, bold=True, spacing=1.8)

tf = txt(s, 1.0, 2.15, 8.2, 1.9)
par(tf, "Website &", size=52, color=INK, bold=True, line=0.92)
par(tf, "Mobile App", size=52, color=GOLD, bold=True, line=0.92)
par(tf, "How to use them", size=24, color=DIM, line=1.0)

rect(s, 1.0, 4.35, 1.4, 0.05, fill=GOLD, rounded=False)

tf = txt(s, 1.0, 4.65, 7.8, 1.0)
par(tf, "A simple, guided walkthrough of everything the website and the app can do "
        "for you and your family — from reels and sermons to giving and prayer.",
    size=14.5, color=DIM, line=1.35)

tf = txt(s, 1.0, 6.15, 8.0, 0.5)
par(tf, [("Presented in church  ·  ", False, DIM), ("Sunday service", True, GOLD_SOFT)],
    size=13, line=1.2)

# QR block on the right
qc = rect(s, 9.95, 2.35, 2.9, 3.0, fill=PANEL, line=LINE)
tf = tframe(qc)
par(tf, "SCAN TO OPEN", size=10.5, color=GOLD, bold=True, align=PP_ALIGN.CENTER, after=6, spacing=1.4)
s.shapes.add_picture(str(ASSETS / "qr_website.png"), Inches(10.35), Inches(2.9), Inches(2.1), Inches(2.1))
tf = txt(s, 9.95, 5.15, 2.9, 0.3, align=PP_ALIGN.CENTER)
par(tf, WEBSITE, size=11, color=GOLD_SOFT, bold=True)

notes(s, """
Welcome and thank the church for the time.
Introduce the session in one line: "Today we are going to learn how to use our church
website and our church app so that nobody misses what God is doing among us."
Keep this short - the goal is a friendly, practical demo, not a technical talk.
If there is a projector and internet, have the website open in a browser tab ready.
""")

# 2 ------------------------------------------------------------ agenda ----
s = slide()
header(s, "Today's session", "What we will cover", "Four short parts — about 20 minutes, then your questions.")
items = [
    ("01", "Why we built this", "One church family, connected every day — not only on Sunday."),
    ("02", "The website", "A guided tour of rccglp63yaya.org.ng, page by page."),
    ("03", "The mobile app", "The CHURCH MEDIA app: reels, Bible, sermons, prayer, alerts."),
    ("04", "Your next steps", "Three simple things to do before you leave today."),
]
for i, (num, title, body) in enumerate(items):
    x = M + i * 3.05
    rect(s, x, 2.55, 2.75, 3.35, fill=PANEL, line=LINE)
    tf = txt(s, x + 0.3, 2.85, 2.2, 0.6)
    par(tf, num, size=30, color=GOLD_DEEP, bold=True)
    tf = txt(s, x + 0.3, 3.55, 2.15, 0.5)
    par(tf, title, size=16, color=INK, bold=True, line=1.05)
    tf = txt(s, x + 0.3, 4.15, 2.15, 1.6)
    par(tf, body, size=12.5, color=DIM, line=1.3)
footer(s, 2)
notes(s, """
Set expectations: four short parts, then questions at the end.
Tell them nothing here requires a login, and nothing here costs money to use.
""")

# 3 --------------------------------------------------------- why it matters
s = slide()
header(s, "Part 1", "Why this matters", "Technology is only a tool — the goal is that nobody is left out of the family.")
cards = [
    ("🔗", "Stay connected", ["Church life continues after Sunday.", "See what God is doing all week long."]),
    ("🎧", "Never miss the Word", ["Catch up on any sermon you missed.", "Listen again and share with a friend."]),
    ("🙏", "Pray & give anywhere", ["Send a prayer request from home.", "Give your tithe and offering securely."]),
    ("🔔", "Always informed", ["Events, changes and announcements.", "Push alerts straight to your phone."]),
]
for i, (ic, t, lines) in enumerate(cards):
    card(s, M + i * 3.05, 2.6, 2.75, 2.5, ic, t, lines)
box = rect(s, M, 5.35, SW - 2 * M, 1.15, fill=RGBColor(0x1C, 0x19, 0x2E), line=RGBColor(0x44, 0x3B, 0x2A))
tf = tframe(box)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
par(tf, "“However far, however near — there's a seat for you.”", size=18, color=GOLD_SOFT, italic=True,
    align=PP_ALIGN.CENTER, line=1.2)
footer(s, 3)
notes(s, """
Emphasise that this is about the people, not the software.
Speak to the families, the youth, the elderly, and those who travel or are unwell -
the website and app let them stay part of the church from anywhere.
""")

# 4 ------------------------------------------------ two ways to connect ---
s = slide()
header(s, "Start here", "Two ways to meet us", "Both show the same church content. Use whichever suits you — or both.")
# website card
rect(s, M, 2.6, 5.55, 3.5, fill=PANEL, line=LINE)
tf = txt(s, M + 0.35, 2.85, 4.9, 0.4)
par(tf, "1.  THE WEBSITE", size=12, color=GOLD, bold=True, spacing=1.4)
tf = txt(s, M + 0.35, 3.28, 4.9, 0.45)
par(tf, WEBSITE, size=22, color=INK, bold=True)
tf = txt(s, M + 0.35, 3.85, 4.9, 1.9)
for i, ln in enumerate(["Works on any phone, tablet or computer.",
                        "Nothing to install — just open and use.",
                        "Best for reading, watching and giving online.",
                        "No login needed for anything public."]):
    par(tf, ln, size=13, color=DIM, line=1.25, after=7, bullet="•", bullet_size=11)
s.shapes.add_picture(str(ASSETS / "qr_website.png"), Inches(5.05), Inches(4.85), Inches(1.05), Inches(1.05))

# app card
rect(s, 6.9, 2.6, 5.55, 3.5, fill=PANEL, line=LINE)
tf = txt(s, 7.25, 2.85, 4.9, 0.4)
par(tf, "2.  THE MOBILE APP", size=12, color=GOLD, bold=True, spacing=1.4)
tf = txt(s, 7.25, 3.28, 4.9, 0.45)
par(tf, "CHURCH MEDIA", size=22, color=INK, bold=True)
tf = txt(s, 7.25, 3.85, 4.9, 1.9)
for ln in ["Free on Google Play — install once.",
           "Reels, offline Bible and audio Bible.",
           "Push notifications for new content.",
           "Works best for daily use on your phone."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=7, bullet="•", bullet_size=11)
s.shapes.add_picture(str(ASSETS / "qr_app.png"), Inches(11.2), Inches(4.85), Inches(1.05), Inches(1.05))
footer(s, 4)
notes(s, """
Make it very practical: the website needs no installation, so anyone can start today.
The app is free on Google Play and gives extra things the website cannot do -
the offline Bible and push notifications.
Point out the two QR codes on screen and say: "Scan either one right now if you like."
""")

# 5 ------------------------------------------------ website getting started
s = slide()
header(s, "Part 2  ·  The website", "Getting to the website — 3 steps", "Nothing to download. Done in under a minute.")
steps = [
    ("1", "Open your browser", "Chrome, Safari, Opera, Edge — whichever is on your phone or computer."),
    ("2", "Type the address", f"Type  {WEBSITE}  in the address bar at the top, then press Go."),
    ("3", "Save it for later", "Tap the menu ⋮ and choose “Add to Home screen”. It becomes a one-tap icon."),
]
for i, (num, t, body) in enumerate(steps):
    y = 2.5 + i * 1.08
    rect(s, M, y, SW - 2 * M, 0.95, fill=PANEL if i % 2 == 0 else PANEL2, line=LINE)
    d = rect(s, M + 0.28, y + 0.19, 0.58, 0.58, fill=GOLD, adj=0.5)
    tf = tframe(d)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    par(tf, num, size=20, color=BG, bold=True, align=PP_ALIGN.CENTER)
    tf = txt(s, M + 1.05, y + 0.2, 3.0, 0.4)
    par(tf, t, size=16, color=INK, bold=True)
    tf = txt(s, M + 4.3, y + 0.2, SW - 2 * M - 4.7, 0.6)
    par(tf, body, size=13, color=DIM, line=1.25)
    if i == 2:
        tf = txt(s, M + 4.3, y + 0.55, SW - 2 * M - 4.7, 0.35)
        par(tf, "Tip: the same trick works for the prayer wall and the giving page.",
            size=11.5, color=GOLD_SOFT, italic=True)
footer(s, 5)
notes(s, """
Walk through this slowly - some members have never typed a web address.
Say the address out loud twice: "r-c-c-g-l-p-6-3-y-a-y-a dot org dot n-g".
The "Add to Home screen" tip is the most valuable thing on this slide -
show it live on your own phone if you can.
""")

# 6 ------------------------------------------------------ website home -----
s = slide()
header(s, "Part 2  ·  The website", "The home page at a glance", "Everything important is one scroll away.")
tf = txt(s, M, 2.45, 5.6, 3.9)
for label, body in [
    ("Welcome banner", "Our welcome message, with buttons to plan a visit or watch the feed."),
    ("From Our Feed", "The newest photos and short reels — swipe through them."),
    ("What's Happening", "Upcoming events with date, time and venue."),
    ("Latest Sermons", "The most recent messages, ready to play."),
    ("Service Times", "Sunday worship, Digging Deep and Faith Clinic."),
    ("Get in Touch / Give", "Contact us or give your tithe and offering in two taps."),
]:
    par(tf, label, size=13.5, color=GOLD_SOFT, bold=True, before=0, after=1, line=1.1)
    par(tf, body, size=12.5, color=DIM, after=8, line=1.2)
picture(s, "home.png", 6.9, 2.45, 5.65)
footer(s, 6)
notes(s, """
If you have internet, open the real website here and scroll it live - it is far more
convincing than a picture.
Point at each section as you name it. Tell them: "If you only remember one thing,
remember that everything starts from this page."
""")

# 7 ----------------------------------------------------- website feed -----
s = slide()
header(s, "Part 2  ·  The website", "Media Feed & Reels", "The life of the church — photos and short videos.")
tf = txt(s, M, 2.5, 5.6, 3.6)
par(tf, "What it is", size=13.5, color=GOLD_SOFT, bold=True, after=3)
for ln in ["A scrolling feed of photos and vertical video “reels”.",
           "Moments from worship, youth nights, outreaches and events."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
par(tf, "How to use it", size=13.5, color=GOLD_SOFT, bold=True, before=10, after=3)
for ln in ["Tap Feed in the top menu.",
           "Scroll down to browse, or tap a reel to open it full screen.",
           "Like, comment and save the ones that bless you.",
           "Use the category chips to filter what you see."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
site_label(s, M, 6.25)
picture(s, "feed.png", 6.9, 2.5, 5.65)
footer(s, 7)
notes(s, """
This is the most engaging part for the youth - appeal to them directly here.
Explain that reels are vertical short videos, like the ones they already watch,
but with content from their own church.
Encourage them: "When you see something good, share it - that is how we reach others."
""")

# 8 --------------------------------------------------- website events -----
s = slide()
header(s, "Part 2  ·  The website", "Events — what's happening", "The church calendar, always up to date.")
tf = txt(s, M, 2.5, 5.6, 3.6)
par(tf, "Every programme in one place", size=13.5, color=GOLD_SOFT, bold=True, after=3)
for ln in ["Upcoming events with date, time and venue.",
           "Tap any event to open its full details and flyer.",
           "Check here before you travel — times can change."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
par(tf, "Who should use it", size=13.5, color=GOLD_SOFT, bold=True, before=10, after=3)
for ln in ["Everyone — especially units, workers and the youth.",
           "Share an event link with a friend who is not yet with us."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
site_label(s, M, 6.25)
picture(s, "events.png", 6.9, 2.5, 5.65)
footer(s, 8)
notes(s, """
Tell them plainly: "If you are not sure whether an event is holding, check the website."
Encourage unit leaders to make sure their programmes are published so the whole
church can see them.
""")

# 9 -------------------------------------------------- website sermons -----
s = slide()
header(s, "Part 2  ·  The website", "Sermons — the Word, on demand", "Missed a service? It is waiting for you here.")
tf = txt(s, M, 2.5, 5.6, 3.6)
par(tf, "What you will find", size=13.5, color=GOLD_SOFT, bold=True, after=3)
for ln in ["Audio and video messages from our services.",
           "The speaker, the date and the scripture reference.",
           "Tap a sermon to play it or read its details."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
par(tf, "Why it helps", size=13.5, color=GOLD_SOFT, bold=True, before=10, after=3)
for ln in ["Re-listen to what God said and take notes.",
           "Send a message to someone going through a hard time.",
           "Perfect for those who were absent, travelling or unwell."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
site_label(s, M, 6.25)
picture(s, "sermons.png", 6.9, 2.5, 5.65)
footer(s, 9)
notes(s, """
Testimony opportunity: if someone was blessed by a message they heard again online,
let them say one sentence here.
Remind them that the same sermons are also inside the app.
""")

# 10 ---------------------------------------------------- website live -----
s = slide()
header(s, "Part 2  ·  The website", "Watch Live", "Join the service from wherever you are.")
tf = txt(s, M, 2.5, 5.6, 3.6)
par(tf, "When we are live", size=13.5, color=GOLD_SOFT, bold=True, after=3)
for ln in ["Open the website and tap “Watch Live Now” on the home page.",
           "The button is shown automatically the moment the stream starts.",
           "Use the Live page any time to see service times."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
par(tf, "Catching up later", size=13.5, color=GOLD_SOFT, bold=True, before=10, after=3)
for ln in ["If you missed it, the message is added to Sermons and Media.",
           "Invite family abroad to join the Sunday service online."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
site_label(s, M, 6.25)
picture(s, "live.png", 6.9, 2.5, 5.65)
footer(s, 10)
notes(s, """
This is a big one for our members abroad and for those who are sick or nursing a baby.
Mention that the stream works on a normal phone internet connection.
Ask the media team to be sure the live flag is switched on before service.
""")

# 11 -------------------------------------------- website bible & search ---
s = slide()
header(s, "Part 2  ·  The website", "Bible & Search", "Read the Word, then find anything on the site.")
card(s, M, 2.55, 5.6, 3.55, "📖", "Bible",
     ["Open Bible from the main menu.",
      "Read the Scriptures right in your browser.",
      "Jump to any book, chapter or verse.",
      "No sign-up, no cost — open and read."])
card(s, 6.9, 2.55, 5.6, 3.55, "🔍", "Search",
     ["Look for anything on the website.",
      "Find a sermon, an event or a post by keyword.",
      "Great when you remember a word but not the title.",
      "Reach it from the footer link on any page."])
footer(s, 11)
notes(s, """
The Bible here is for quick reading on any device.
For daily reading and listening, the app is better because the Bible works without
internet and can read aloud to you - that is on the app slides coming next.
""")

# 12 --------------------------------------------------- website prayer -----
s = slide()
header(s, "Part 2  ·  The website", "Prayer Wall", "We pray for one another — from anywhere in the world.")
card(s, M, 2.55, 5.6, 3.55, "🙏", "Submit a request",
     ["Open Prayer from the menu or the footer.",
      "Type your request and send it in.",
      "Choose to keep it private if it is personal.",
      "You do not need an account to send one."])
card(s, 6.9, 2.55, 5.6, 3.55, "❤️", "Pray for others",
     ["Read the requests others have shared.",
      "Stand in agreement with the church family.",
      "Use it with your unit, family or prayer group.",
      "Prayer requests are monitored by the prayer team."])
footer(s, 12)
notes(s, """
Invite someone to type a request today as a demonstration.
Remind them that private requests are seen by the pastoral/prayer team only.
This is a strong point for those who find it hard to speak out in person.
""")

# 13 ---------------------------------------------------- website give -----
s = slide()
header(s, "Part 2  ·  The website", "Giving Online", "Tithes, offerings and pledges — safely and securely.")
tf = txt(s, M, 2.5, 5.6, 3.6)
par(tf, "How to give", size=13.5, color=GOLD_SOFT, bold=True, after=3)
for ln in ["Tap “Give” in the menu or “Give Online” on the home page.",
           "Choose what you are giving — tithe, offering, pledge or project.",
           "Pay by card or bank transfer through our secure payment channel.",
           "You receive a confirmation for your records."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
par(tf, "Please note", size=13.5, color=GOLD_SOFT, bold=True, before=10, after=3)
for ln in ["Give only through the official links on our website or app.",
           "Never send money to a personal account that is not ours."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
site_label(s, M, 6.25)
picture(s, "give.png", 6.9, 2.5, 5.65)
footer(s, 13)
notes(s, """
Be clear and reassuring: this is the same church account, just a safer and easier way
to give. Mention that records are kept properly by the finance team.
Strongly warn against fraudsters who ask for money into personal accounts.
""")

# 14 --------------------------------------------------- website units -----
s = slide()
header(s, "Part 2  ·  The website", "Find your parish", "Our network: Province → Zone → Area → Parish.")
tf = txt(s, M, 2.5, 5.6, 3.6)
par(tf, "Browse the network", size=13.5, color=GOLD_SOFT, bold=True, after=3)
for ln in ["Open “Parishes” in the main menu.",
           "Browse from province down to area and parish level.",
           "Select a parish to see its details.",
           "Ideal for members who relocate or travel."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
par(tf, "Good for", size=13.5, color=GOLD_SOFT, bold=True, before=10, after=3)
for ln in ["Finding the nearest parish when you travel.",
           "Inviting someone who lives far from us to their local parish.",
           "Workers who need to see how our network is arranged."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
site_label(s, M, 6.25)
picture(s, "units.png", 6.9, 2.5, 5.65)
footer(s, 14)
notes(s, """
Use this to help anyone who travels often or has family elsewhere.
Explain the structure simply: province, then zone, then area, then parish -
so they know where they belong and where they can fellowship.
""")

# 15 ------------------------------------------- website about & contact ---
s = slide()
header(s, "Part 2  ·  The website", "About, Contact & Updates", "Who we are, how to reach us, and how to hear first.")
card(s, M, 2.55, 3.55, 3.55, "ℹ️", "About Us",
     ["Our history and message.",
      "What we believe and where we are going.",
      "A good page to send a first-time visitor.",
      "Reached from the About menu."], title_size=14, body_size=12)
card(s, 4.9, 2.55, 3.55, 3.55, "📞", "Contact",
     ["Our phone numbers and email.",
      "Reach the church office directly.",
      "Ask a question or plan a visit.",
      "Perfect for first-time guests."], title_size=14, body_size=12)
card(s, 9.05, 2.55, 3.55, 3.55, "✉️", "Newsletter",
     ["Enter your email in the footer box.",
      "Get updates and announcements by email.",
      "One click to join — no app needed.",
      "Great for parents and workers."], title_size=14, body_size=12)
footer(s, 15)
notes(s, """
Show the contact details out loud - someone may want to note them down.
Encourage the congregation to help a first-time visitor by sharing the website.
Mention the newsletter as an option for those who prefer email over the app.
""")

# 16 --------------------------------------------------- advertise with us -
s = slide()
header(s, "Part 2  ·  The website", "Advertise with us", "A door for businesses and ministries in our church family.")
tf = txt(s, M, 2.5, 5.6, 3.6)
par(tf, "For business owners and ministries", size=13.5, color=GOLD_SOFT, bold=True, after=3)
for ln in ["Show your product, brand or ministry to our church audience.",
           "Choose to appear on the website, in the app, or both.",
           "Send a vertical (9:16) image or short video advert.",
           "Free trial slots are available — a paid premium option exists."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
par(tf, "Already advertising?", size=13.5, color=GOLD_SOFT, bold=True, before=10, after=3)
for ln in ["Use the Publisher Portal to track views and clicks.",
           "Request an access link with your email address."]:
    par(tf, ln, size=13, color=DIM, line=1.25, after=6, bullet="•", bullet_size=11)
site_label(s, M, 6.25)
picture(s, "advertise.png", 6.9, 2.5, 5.65)
footer(s, 16)
notes(s, """
This slide may not apply to everyone - keep it brief.
Speak directly to entrepreneurs, traders and ministry leaders in the house.
Remind them that adverts must be approved by the church before they go live.
""")

# 17 ------------------------------------------------------- app download --
s = slide()
header(s, "Part 3  ·  The mobile app", "Getting the app — 4 steps", "Free on Google Play. Takes two minutes.")
for i, (num, t, body) in enumerate([
    ("1", "Open the Play Store", "On your Android phone, tap the Google Play Store app."),
    ("2", "Find the app", "Search “CHURCH MEDIA”, or scan the QR code on this slide."),
    ("3", "Tap Install", "It is 100% free. Wait for it to finish, then tap Open."),
    ("4", "Allow notifications", "Say “Allow” so you receive our announcements."),
]):
    x = M + (i % 2) * 6.15
    y = 2.55 + (i // 2) * 1.7
    rect(s, x, y, 5.6, 1.45, fill=PANEL if i % 2 == 0 else PANEL2, line=LINE)
    d = rect(s, x + 0.25, y + 0.35, 0.6, 0.6, fill=GOLD, adj=0.5)
    tf = tframe(d)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    par(tf, num, size=20, color=BG, bold=True, align=PP_ALIGN.CENTER)
    tf = txt(s, x + 1.05, y + 0.3, 4.3, 0.4)
    par(tf, t, size=15, color=INK, bold=True)
    tf = txt(s, x + 1.05, y + 0.72, 4.3, 0.6)
    par(tf, body, size=12.5, color=DIM, line=1.25)

s.shapes.add_picture(str(ASSETS / "qr_app.png"), Inches(9.55), Inches(5.85), Inches(1.15), Inches(1.15))
tf = txt(s, 10.85, 6.05, 2.1, 0.8)
par(tf, "Scan to install", size=12, color=GOLD_SOFT, bold=True, line=1.15)
par(tf, "Search “CHURCH MEDIA”", size=11, color=DIM, line=1.15)
footer(s, 17)
notes(s, """
Demo this on your own phone and pass it round if possible.
Say clearly: the app is free, and it is the same church - not a third-party company.
If someone has an iPhone, tell them the website works fully on iPhone and the app is
currently on Google Play for Android.
""")

# 18 -------------------------------------------------------- five tabs -----
s = slide()
header(s, "Part 3  ·  The mobile app", "Inside the app — five tabs", "The bar at the bottom of the app takes you everywhere.")
# phone mock-up
px, py, pw, ph = 1.15, 2.45, 2.5, 4.25
rect(s, px, py, pw, ph, fill=RGBColor(0x26, 0x22, 0x40), line=GOLD_DEEP, lw=1.5, adj=0.09)
rect(s, px + 0.13, py + 0.13, pw - 0.26, ph - 0.26, fill=RGBColor(0x10, 0x0F, 0x1E), line=None, adj=0.07)
tf = txt(s, px + 0.3, py + 0.35, pw - 0.6, 0.9)
par(tf, "CHURCH", size=13, color=GOLD, bold=True, align=PP_ALIGN.CENTER, line=1.0)
par(tf, "MEDIA", size=13, color=GOLD, bold=True, align=PP_ALIGN.CENTER, line=1.0)
par(tf, "RCCG LP63 YAYA", size=8.5, color=DIM, align=PP_ALIGN.CENTER, before=3)
rect(s, px + 0.32, py + 1.42, pw - 0.64, 0.9, fill=RGBColor(0x1D, 0x1A, 0x30), line=LINE, adj=0.1)
rect(s, px + 0.32, py + 2.44, pw - 0.64, 0.9, fill=RGBColor(0x1D, 0x1A, 0x30), line=LINE, adj=0.1)
tf = txt(s, px + 0.46, py + 1.62, pw - 0.92, 0.6)
par(tf, "Latest sermons", size=10.5, color=DIM, line=1.1)
tf = txt(s, px + 0.46, py + 2.64, pw - 0.92, 0.6)
par(tf, "Upcoming events", size=10.5, color=DIM, line=1.1)
nav = ["Home", "Feed", "Bible", "Sermons", "More"]
for i, n in enumerate(nav):
    tf = txt(s, px + 0.1 + i * (pw - 0.2) / 5, py + ph - 0.52, (pw - 0.2) / 5, 0.32, align=PP_ALIGN.CENTER)
    par(tf, n, size=7.5, color=GOLD_SOFT if i == 0 else DIM, bold=i == 0)

tabs = [
    ("Home", "Service times, latest sermons and quick shortcuts."),
    ("Feed", "Full-screen reels and photos — swipe to browse."),
    ("Bible", "Read and listen, even without internet."),
    ("Sermons", "The full message archive on audio and video."),
    ("More", "Events, Live, Prayer, About, Contact, Give and Search."),
]
for i, (name, desc) in enumerate(tabs):
    y = 2.55 + i * 0.78
    rect(s, 4.35, y, 8.15, 0.68, fill=PANEL if i % 2 == 0 else PANEL2, line=LINE)
    tf = txt(s, 4.6, y + 0.19, 1.7, 0.35)
    par(tf, name, size=14, color=GOLD_SOFT, bold=True)
    tf = txt(s, 6.35, y + 0.2, 5.9, 0.4)
    par(tf, desc, size=12.5, color=DIM, line=1.15)
footer(s, 18)
notes(s, """
Say it simply: "Five buttons at the bottom - that is the whole app."
Home first, then Feed for reels, Bible for the Word, Sermons for messages,
and More for everything else.
Tell them: "If you ever feel lost, just tap Home."
""")

# 19 ---------------------------------------------------------- app reels --
s = slide()
header(s, "Part 3  ·  The app", "Reels & Media Feed", "Church moments in the format young people already love.")
card(s, M, 2.55, 3.85, 3.55, "🎬", "Watch",
     ["Swipe up and down for the next reel.",
      "Full-screen, vertical, fast.",
      "Pinned reels appear first.",
      "Photos and videos together."], title_size=14.5, body_size=12.5)
card(s, 4.95, 2.55, 3.85, 3.55, "❤️", "React",
     ["Like the reels that bless you.",
      "Save them to find later.",
      "Comment and reply, Instagram-style.",
      "See what others are saying."], title_size=14.5, body_size=12.5)
card(s, 9.05, 2.55, 3.55, 3.55, "🏷️", "Filter",
     ["Use categories to sort content.",
      "Worship, sermon clips, youth and events.",
      "Great for youth and young adults.",
      "New content appears automatically."], title_size=14.5, body_size=12.5)
footer(s, 19)
notes(s, """
This is the feature that will excite the youth the most.
Invite a young person to demonstrate on their phone for 30 seconds if you can.
Encourage them to follow the church and share reels with their friends.
""")

# 20 ----------------------------------------------------------- app bible --
s = slide()
header(s, "Part 3  ·  The app", "Bible — read and listen, even offline", "The strongest reason to install the app.")
card(s, M, 2.55, 2.7, 3.55, "📶", "Offline",
     ["KJV and BBE read fully offline.",
      "No data, no airtime needed.",
      "Perfect in church or while travelling.",
      "Always available on your phone."], title_size=14, body_size=12)
card(s, 3.8, 2.55, 2.7, 3.55, "🔊", "Listen",
     ["Audio Bible reads to you.",
      "The verse being read is highlighted.",
      "Speed: 0.75x, 1x, 1.25x, 1.5x.",
      "Auto-scroll follows the reading."], title_size=14, body_size=12)
card(s, 6.85, 2.55, 2.7, 3.55, "🌍", "Versions",
     ["NIV, NLT and NKJV online.",
      "Multi-language support.",
      "Switch version in one tap.",
      "Ideal for study and devotion."], title_size=14, body_size=12)
card(s, 9.9, 2.55, 2.7, 3.55, "🔎", "Search",
     ["Search any verse instantly.",
      "The page scrolls straight to it.",
      "No more flipping endlessly.",
      "Use it during service."], title_size=14, body_size=12)
footer(s, 20)
notes(s, """
Challenge the church: "Let us read the Bible every day this week."
The offline KJV is a real blessing for those who cannot afford data every day.
Show the audio Bible on your own phone - let them hear it read a verse aloud.
""")

# 21 ------------------------------------------- app sermons/events/live ---
s = slide()
header(s, "Part 3  ·  The app", "Sermons, Events & Live", "Everything you would come to church for — in your pocket.")
card(s, M, 2.55, 3.85, 3.55, "🎧", "Sermons",
     ["Full sermon archive on audio and video.",
      "Speaker, date and scripture included.",
      "Replay and share any message.",
      "Build your own daily listening habit."], title_size=14.5, body_size=12.5)
card(s, 4.95, 2.55, 3.85, 3.55, "🗓️", "Events",
     ["Upcoming programmes with full details.",
      "Tap an event for the venue and time.",
      "Plan ahead with your family.",
      "Never hear about an event too late."], title_size=14.5, body_size=12.5)
card(s, 9.05, 2.55, 3.55, 3.55, "📺", "Live",
     ["Watch the service live from the app.",
      "Ideal when you cannot attend.",
      "Share with family far away.",
      "Catch up later in Sermons."], title_size=14.5, body_size=12.5)
footer(s, 21)
notes(s, """
Connect this to real people: the student away at school, the nursing mother,
the traveller, the one who is unwell.
Say: "Church is not only a building - it is a family, and this keeps you in the family."
""")

# 22 ---------------------------------------- app prayer/give/notifications -
s = slide()
header(s, "Part 3  ·  The app", "Prayer, Giving & Notifications", "Stay connected in the moments that matter.")
card(s, M, 2.55, 3.85, 3.55, "🙏", "Prayer Wall",
     ["Send a prayer request from your phone.",
      "Public or private - your choice.",
      "Pray for others on the wall.",
      "Open from More → Prayer Wall."], title_size=14.5, body_size=12.5)
card(s, 4.95, 2.55, 3.85, 3.55, "💛", "Give",
     ["Tithes, offerings and pledges.",
      "Secure payment in a few taps.",
      "Open from More → Give.",
      "Confirmation for your records."], title_size=14.5, body_size=12.5)
card(s, 9.05, 2.55, 3.55, 3.55, "🔔", "Notifications",
     ["Alerts for new posts and events.",
      "Announcements reach you first.",
      "Tap an alert to open the content.",
      "Turn them on in your phone settings."], title_size=14.5, body_size=12.5)
footer(s, 22)
notes(s, """
The notification point is important - many people install the app but never allow
notifications, then wonder why they hear nothing.
Ask them to check it today: Phone Settings → Apps → CHURCH MEDIA → Notifications → Allow.
""")

# 23 ---------------------------------------------------------- reference --
s = slide()
header(s, "Quick reference", "“Where do I go for…?”", "Keep this slide — it answers almost every question.")
table(s, M, 2.4, SW - 2 * M, 4.3, [
    ["I want to…", "On the website", "In the app"],
    ["Watch short videos & reels", "Feed", "Feed tab"],
    ["Read the Bible offline", "Bible (online only)", "Bible tab — KJV/BBE offline"],
    ["Read or listen to the Bible", "Bible", "Bible tab — with audio"],
    ["Hear a past sermon", "Sermons", "Sermons tab"],
    ["See upcoming programmes", "Events", "More → Events"],
    ["Watch the service live", "Live", "More → Live"],
    ["Send a prayer request", "Prayer", "More → Prayer Wall"],
    ["Give my tithe or offering", "Give", "More → Give"],
    ["Contact the church", "Contact", "More → Contact"],
], widths=[3.3, 2.6, 3.2], row_h=0.4)
footer(s, 23)
notes(s, """
Give them a moment to photograph this slide.
Say: "Everything on this table is free and needs no login, except giving where you
choose your own payment details."
""")

# 24 ------------------------------------------------------ call to action --
s = slide()
header(s, "Your next steps", "Three things to do this week", "Small habits that keep you connected to the family.")
acts = [
    ("1", "Save the website", f"Open {WEBSITE} and add it to your home screen."),
    ("2", "Install the app", "Search “CHURCH MEDIA” on Google Play and tap Install."),
    ("3", "Turn on notifications — and invite someone",
     "Allow alerts, then share the website with one person who was not here today."),
]
for i, (num, t, body) in enumerate(acts):
    y = 2.55 + i * 1.12
    rect(s, M, y, 8.15, 1.0, fill=PANEL if i % 2 == 0 else PANEL2, line=LINE)
    d = rect(s, M + 0.28, y + 0.21, 0.58, 0.58, fill=GOLD, adj=0.5)
    tf = tframe(d)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    par(tf, num, size=20, color=BG, bold=True, align=PP_ALIGN.CENTER)
    tf = txt(s, M + 1.05, y + 0.16, 6.9, 0.4)
    par(tf, t, size=15.5, color=INK, bold=True)
    tf = txt(s, M + 1.05, y + 0.55, 6.9, 0.4)
    par(tf, body, size=12.5, color=DIM, line=1.2)

rect(s, 9.4, 2.55, 3.2, 3.56, fill=PANEL, line=LINE)
tf = txt(s, 9.4, 2.78, 3.2, 0.3, align=PP_ALIGN.CENTER)
par(tf, "SCAN & GO", size=11, color=GOLD, bold=True, spacing=1.4, align=PP_ALIGN.CENTER)
s.shapes.add_picture(str(ASSETS / "qr_website.png"), Inches(9.62), Inches(3.22), Inches(1.32), Inches(1.32))
s.shapes.add_picture(str(ASSETS / "qr_app.png"), Inches(11.06), Inches(3.22), Inches(1.32), Inches(1.32))
tf = txt(s, 9.62, 4.62, 1.32, 0.3, align=PP_ALIGN.CENTER)
par(tf, "Website", size=10.5, color=GOLD_SOFT, bold=True, align=PP_ALIGN.CENTER)
tf = txt(s, 11.06, 4.62, 1.32, 0.3, align=PP_ALIGN.CENTER)
par(tf, "The app", size=10.5, color=GOLD_SOFT, bold=True, align=PP_ALIGN.CENTER)
tf = txt(s, 9.55, 5.06, 2.9, 0.9, align=PP_ALIGN.CENTER)
par(tf, "Scan one now and follow along with us this morning.", size=11, color=DIM,
    align=PP_ALIGN.CENTER, line=1.25)
footer(s, 24)
notes(s, """
End with a clear, simple ask - three actions, nothing more.
If the church has a projector and internet, ask everyone who can to scan the QR code now
and install the app before they leave the hall.
""")

# 25 -------------------------------------------------------- thank you ----
s = slide(page=False)
rect(s, 0, 0, SW, 0.14, fill=GOLD_DEEP, rounded=False)
tf = txt(s, 1.0, 1.5, 11.3, 1.6, align=PP_ALIGN.CENTER)
par(tf, "Thank you", size=46, color=INK, bold=True, align=PP_ALIGN.CENTER, line=1.0)
par(tf, "Questions & answers", size=20, color=GOLD, align=PP_ALIGN.CENTER, before=6)

rect(s, 1.0, 3.55, 11.33, 0.02, fill=LINE, rounded=False)

tf = txt(s, 1.0, 3.85, 11.33, 1.6, align=PP_ALIGN.CENTER)
par(tf, "“However far, however near — there's a seat for you.”",
    size=19, color=GOLD_SOFT, italic=True, align=PP_ALIGN.CENTER, after=14)
par(tf, f"{WEBSITE}   ·   CHURCH MEDIA on Google Play", size=14.5, color=INK,
    bold=True, align=PP_ALIGN.CENTER)
par(tf, "admin@rccglp63yaya.org.ng", size=13, color=DIM, align=PP_ALIGN.CENTER, before=6)
par(tf, PHONES, size=12, color=DIM, align=PP_ALIGN.CENTER, before=3)

rect(s, 1.0, 6.15, 11.33, 0.75, fill=PANEL, line=LINE)
tf = tframe(s.shapes[-1])
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
par(tf, "Sunday Worship 9:00–11:30 AM   ·   Digging Deep (Tue) 5:00–6:00 PM   ·   Faith Clinic 5:00–6:00 PM",
    size=12.5, color=DIM, align=PP_ALIGN.CENTER)

notes(s, """
Open the floor for questions. Have the website open on a phone or laptop so you can
answer live - that is the most powerful moment of the whole session.
If nobody asks, ask them: "Who has already installed the app? Please raise your hand."
Close with prayer.
""")

# ---------------------------------------------------------------- meta ----
prs.core_properties.title = "RCCG LP63 YAYA — Website & Mobile App Guide"
prs.core_properties.author = "RCCG LP63 YAYA Media Team"
prs.core_properties.subject = "How to use the church website and mobile app"
prs.core_properties.comments = "Presented in church. Website: rccglp63yaya.org.ng"

prs.save(OUT)
print(f"Saved: {OUT}")
print(f"Slides: {len(prs.slides._sldIdLst)}")
