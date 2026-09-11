#!/usr/bin/env python3
"""
Generates Church_Media_Presentation.pptx for the upcoming Sunday Church Meeting.
Uses python-pptx to build a 10-slide, 16:9 widescreen presentation deck with a
dark/gold design system matching the Church Media website & mobile app.
"""

import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    COLOR_BG = RGBColor(10, 9, 18)         # #0A0912
    COLOR_PANEL = RGBColor(22, 19, 38)     # #161326
    COLOR_BORDER = RGBColor(44, 40, 80)    # #2C2850
    COLOR_GOLD = RGBColor(232, 185, 95)    # #E8B95F
    COLOR_GOLD_SOFT = RGBColor(243, 211, 143) # #F3D38F
    COLOR_WHITE = RGBColor(255, 255, 255)  # #FFFFFF
    COLOR_DIM = RGBColor(220, 216, 242)    # #DCD8F2
    COLOR_FAINT = RGBColor(155, 150, 188)  # #9B96BC

    blank_layout = prs.slide_layouts[6]

    def set_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()

    def add_header(slide, category, title):
        # Category Eyebrow
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(12)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_GOLD_SOFT

        # Main Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.733), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)

    # Decorative Card in Center
    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1))
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_PANEL
    card1.line.color.rgb = COLOR_GOLD
    card1.line.width = Pt(1.5)

    tb1 = s1.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(9.733), Inches(3.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "CHURCH MEDIA MANAGEMENT SYSTEM"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_GOLD_SOFT
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf1.add_paragraph()
    p2.text = "Connecting Our Ministry Across Website & Mobile App"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(16)

    p3 = tf1.add_paragraph()
    p3.text = "Official Presentation for Church Service & Meeting"
    p3.font.size = Pt(18)
    p3.font.color.rgb = COLOR_DIM
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(20)

    p4 = tf1.add_paragraph()
    p4.text = "Available on Google Play & Web • Full Reel Feed • Offline Bible • Live Stream • Online Giving"
    p4.font.size = Pt(13)
    p4.font.color.rgb = COLOR_FAINT
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(24)

    # -------------------------------------------------------------
    # SLIDE 2: Vision & Executive Summary
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "Overview & Vision", "A Digital Ministry Without Walls")

    left_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = COLOR_PANEL
    left_box.line.color.rgb = COLOR_BORDER

    tf_l = left_box.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.3)
    tf_l.margin_right = Inches(0.3)
    tf_l.margin_top = Inches(0.3)

    pl1 = tf_l.paragraphs[0]
    pl1.text = "🎯 Why We Built This System"
    pl1.font.size = Pt(20)
    pl1.font.bold = True
    pl1.font.color.rgb = COLOR_GOLD

    bullets_l = [
        ("Unified Digital Hub:", "One platform connecting our website, Android Mobile App, and parish network."),
        ("Reaching Every Generation:", "Engaging youth, families, and global members through interactive media reels."),
        ("Seamless Accessibility:", "Instant access to live broadcasts, sermon audio, events, and offline scripture."),
        ("Parish & Province Network:", "Multi-unit support connecting Province, Zone, Area, and Parish churches.")
    ]
    for b_title, b_desc in bullets_l:
        p = tf_l.add_paragraph()
        p.space_before = Pt(12)
        run1 = p.add_run()
        run1.text = "• " + b_title + " "
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = COLOR_GOLD_SOFT
        run2 = p.add_run()
        run2.text = b_desc
        run2.font.size = Pt(13)
        run2.font.color.rgb = COLOR_DIM

    right_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.8))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = COLOR_PANEL
    right_box.line.color.rgb = COLOR_BORDER

    tf_r = right_box.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.3)
    tf_r.margin_right = Inches(0.3)
    tf_r.margin_top = Inches(0.3)

    pr1 = tf_r.paragraphs[0]
    pr1.text = "📱 Website & Mobile App Highlights"
    pr1.font.size = Pt(20)
    pr1.font.bold = True
    pr1.font.color.rgb = COLOR_GOLD

    bullets_r = [
        ("Mobile App on Google Play:", "Download directly on Android phones for notifications & offline Bible."),
        ("Full-Screen Vertical Reels:", "Swipeable 9:16 reels for worship, sermons, and testimony clips."),
        ("Voice-Assisted Audio Bible:", "Listen to scripture with verse-by-verse audio and active highlighting."),
        ("Online Giving & Ad Manager:", "Pay tithes/offerings online or place vertical ads to reach our audience.")
    ]
    for b_title, b_desc in bullets_r:
        p = tf_r.add_paragraph()
        p.space_before = Pt(12)
        run1 = p.add_run()
        run1.text = "• " + b_title + " "
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = COLOR_GOLD_SOFT
        run2 = p.add_run()
        run2.text = b_desc
        run2.font.size = Pt(13)
        run2.font.color.rgb = COLOR_DIM

    # -------------------------------------------------------------
    # SLIDE 3: Reels & Media Feed
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "Core Experience", "Full-Screen Media Reels & Feed")

    # 3 Cards Layout
    card_w = Inches(3.64)
    card_h = Inches(4.8)

    cards_data3 = [
        ("🎬 Swipeable 9:16 Reels", [
            ("Vertical Video Format:", "Designed for immersive mobile viewing just like Reels & TikTok."),
            ("Category Filtering:", "Browse by Worship, Sermon Clip, Youth, Testimony, or Events."),
            ("Instant Streaming:", "Videos play instantly on both web and mobile app.")
        ]),
        ("💬 Interactive Engagement", [
            ("Likes & Saves:", "Visitors can like and bookmark favorite ministry moments."),
            ("Instagram-Style Comments:", "Threaded comment replies with emoji picker and image attachments."),
            ("Community Shield:", "Rate-limited submission prevents spam.")
        ]),
        ("📌 Pinned Featured Reels", [
            ("Top Priority:", "Church leaders can pin up to 3 featured reels per church."),
            ("Auto-Expiry:", "Pinned posts stay featured for 3 days automatically."),
            ("Multi-Level Rollup:", "Pins surface on local Parish, Zone, Province, and main Feed.")
        ])
    ]

    for i, (ctitle, citems) in enumerate(cards_data3):
        left_pos = Inches(0.8 + i * 3.95)
        cbox = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.8), card_w, card_h)
        cbox.fill.solid()
        cbox.fill.fore_color.rgb = COLOR_PANEL
        cbox.line.color.rgb = COLOR_BORDER

        tf = cbox.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = ctitle
        p0.font.size = Pt(18)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_GOLD

        for head, body in citems:
            p = tf.add_paragraph()
            p.space_before = Pt(12)
            r1 = p.add_run()
            r1.text = "• " + head + " "
            r1.font.bold = True
            r1.font.size = Pt(12)
            r1.font.color.rgb = COLOR_GOLD_SOFT
            r2 = p.add_run()
            r2.text = body
            r2.font.size = Pt(12)
            r2.font.color.rgb = COLOR_DIM

    # -------------------------------------------------------------
    # SLIDE 4: Holy Bible & Audio Reader
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "Spiritual Growth", "Holy Bible & Voice-Assisted Audio Reader")

    b_left = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    b_left.fill.solid()
    b_left.fill.fore_color.rgb = COLOR_PANEL
    b_left.line.color.rgb = COLOR_BORDER
    tf_bl = b_left.text_frame
    tf_bl.word_wrap = True
    tf_bl.margin_left = Inches(0.3)
    tf_bl.margin_right = Inches(0.3)
    tf_bl.margin_top = Inches(0.3)

    pbl0 = tf_bl.paragraphs[0]
    pbl0.text = "🔊 Voice Reading & TTS Audio Player"
    pbl0.font.size = Pt(20)
    pbl0.font.bold = True
    pbl0.font.color.rgb = COLOR_GOLD

    items_bl = [
        ("Listen to Any Chapter:", "Integrated Web Speech API reads scripture aloud with natural voice synthesis."),
        ("Active Verse Highlighting:", "Highlights each verse in real-time as it is spoken."),
        ("Auto-Scrolling:", "Automatically scrolls down the chapter so you can read along effortlessly."),
        ("Speed Control:", "Adjust playback speed from 0.75x, 1.0x, 1.25x to 1.5x for personal study.")
    ]
    for head, desc in items_bl:
        p = tf_bl.add_paragraph()
        p.space_before = Pt(12)
        r1 = p.add_run()
        r1.text = "• " + head + " "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_GOLD_SOFT
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_DIM

    b_right = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.8))
    b_right.fill.solid()
    b_right.fill.fore_color.rgb = COLOR_PANEL
    b_right.line.color.rgb = COLOR_BORDER
    tf_br = b_right.text_frame
    tf_br.word_wrap = True
    tf_br.margin_left = Inches(0.3)
    tf_br.margin_right = Inches(0.3)
    tf_br.margin_top = Inches(0.3)

    pbr0 = tf_br.paragraphs[0]
    pbr0.text = "📚 Translations & Offline Study"
    pbr0.font.size = Pt(20)
    pbr0.font.bold = True
    pbr0.font.color.rgb = COLOR_GOLD

    items_br = [
        ("100% Offline Versions:", "King James Version (KJV) and Bible in Basic English (BBE) work completely offline on mobile."),
        ("Modern Translations:", "Access NIV, NLT, NKJV, and multi-language translations online."),
        ("Verse Search & Jump:", "Type a verse or keyword and jump straight to the exact passage."),
        ("Chapter Navigation:", "Easily switch books, chapters, and testaments with one tap.")
    ]
    for head, desc in items_br:
        p = tf_br.add_paragraph()
        p.space_before = Pt(12)
        r1 = p.add_run()
        r1.text = "• " + head + " "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_GOLD_SOFT
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_DIM

    # -------------------------------------------------------------
    # SLIDE 5: G.O. Declaration & Livestream
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "Global Leadership & Broadcasts", "G.O. Declaration & Live Services")

    card_s5_left = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    card_s5_left.fill.solid()
    card_s5_left.fill.fore_color.rgb = COLOR_PANEL
    card_s5_left.line.color.rgb = COLOR_GOLD

    tf_s5l = card_s5_left.text_frame
    tf_s5l.word_wrap = True
    tf_s5l.margin_left = Inches(0.3)
    tf_s5l.margin_right = Inches(0.3)
    tf_s5l.margin_top = Inches(0.3)

    ps5l0 = tf_s5l.paragraphs[0]
    ps5l0.text = "✨ General Overseer (G.O.) Declaration"
    ps5l0.font.size = Pt(20)
    ps5l0.font.bold = True
    ps5l0.font.color.rgb = COLOR_GOLD

    items_s5l = [
        ("Prophetic Word Banner:", "Displays the General Overseer's theme and annual prophetic declaration prominently."),
        ("Marquee or Static Mode:", "Super admin can choose an animated scrolling marquee or a bold fixed top banner."),
        ("Site-Wide Visibility:", "Positioned at the top of every page so members never miss spiritual direction.")
    ]
    for head, desc in items_s5l:
        p = tf_s5l.add_paragraph()
        p.space_before = Pt(14)
        r1 = p.add_run()
        r1.text = "• " + head + " "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_GOLD_SOFT
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_DIM

    card_s5_right = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.8))
    card_s5_right.fill.solid()
    card_s5_right.fill.fore_color.rgb = COLOR_PANEL
    card_s5_right.line.color.rgb = COLOR_BORDER

    tf_s5r = card_s5_right.text_frame
    tf_s5r.word_wrap = True
    tf_s5r.margin_left = Inches(0.3)
    tf_s5r.margin_right = Inches(0.3)
    tf_s5r.margin_top = Inches(0.3)

    ps5r0 = tf_s5r.paragraphs[0]
    ps5r0.text = "📹 Live Broadcasts & Sermon Archive"
    ps5r0.font.size = Pt(20)
    ps5r0.font.bold = True
    ps5r0.font.color.rgb = COLOR_GOLD

    items_s5r = [
        ("Live Stream Badge:", "Displays a pulsing 'LIVE' badge when Sunday or midweek services are broadcasting."),
        ("Embedded YouTube Live:", "Watch live service broadcasts directly inside the website or mobile app."),
        ("Sermon Library:", "Browse past sermon recordings, audio mp3 downloads, speaker details, and scripture refs.")
    ]
    for head, desc in items_s5r:
        p = tf_s5r.add_paragraph()
        p.space_before = Pt(14)
        r1 = p.add_run()
        r1.text = "• " + head + " "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_GOLD_SOFT
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_DIM

    # -------------------------------------------------------------
    # SLIDE 6: Online Giving & Parish Network
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "Generosity & Community", "Online Giving & Parish Directory")

    card_s6_l = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    card_s6_l.fill.solid()
    card_s6_l.fill.fore_color.rgb = COLOR_PANEL
    card_s6_l.line.color.rgb = COLOR_BORDER

    tf_s6l = card_s6_l.text_frame
    tf_s6l.word_wrap = True
    tf_s6l.margin_left = Inches(0.3)
    tf_s6l.margin_right = Inches(0.3)
    tf_s6l.margin_top = Inches(0.3)

    ps6l0 = tf_s6l.paragraphs[0]
    ps6l0.text = "💳 Secure Online Giving & Tithes"
    ps6l0.font.size = Pt(20)
    ps6l0.font.bold = True
    ps6l0.font.color.rgb = COLOR_GOLD

    items_s6l = [
        ("Multiple Giving Categories:", "Tithes, Offerings, Building Fund, Special Seed, and Missions."),
        ("Payhub Gateway Integration:", "Instant, secure online card & bank checkout."),
        ("Manual Bank Transfer Option:", "Upload proof of bank transfer receipt for verification by finance team."),
        ("Financial Audit Reports:", "Admin dashboard for tracking completed, pending, and verified donations.")
    ]
    for head, desc in items_s6l:
        p = tf_s6l.add_paragraph()
        p.space_before = Pt(12)
        r1 = p.add_run()
        r1.text = "• " + head + " "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_GOLD_SOFT
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_DIM

    card_s6_r = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.8))
    card_s6_r.fill.solid()
    card_s6_r.fill.fore_color.rgb = COLOR_PANEL
    card_s6_r.line.color.rgb = COLOR_BORDER

    tf_s6r = card_s6_r.text_frame
    tf_s6r.word_wrap = True
    tf_s6r.margin_left = Inches(0.3)
    tf_s6r.margin_right = Inches(0.3)
    tf_s6r.margin_top = Inches(0.3)

    ps6r0 = tf_s6r.paragraphs[0]
    ps6r0.text = "🏛️ Multi-Unit Church Directory"
    ps6r0.font.size = Pt(20)
    ps6r0.font.bold = True
    ps6r0.font.color.rgb = COLOR_GOLD

    items_s6r = [
        ("Organisational Hierarchy:", "Province → Zone → Area → Parish structure."),
        ("Dedicated Parish Pages:", "Each parish has its own page, service schedule, leadership team, and media."),
        ("Parish Registration:", "New parishes can register online for super-admin approval and corporate email setup.")
    ]
    for head, desc in items_s6r:
        p = tf_s6r.add_paragraph()
        p.space_before = Pt(14)
        r1 = p.add_run()
        r1.text = "• " + head + " "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_GOLD_SOFT
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_DIM

    # -------------------------------------------------------------
    # SLIDE 7: Advertising System & Publisher Portal
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "Monetization & Outreach", "9:16 Vertical Advertising System")

    card_s7_1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    card_s7_1.fill.solid()
    card_s7_1.fill.fore_color.rgb = COLOR_PANEL
    card_s7_1.line.color.rgb = COLOR_BORDER

    tf_s71 = card_s7_1.text_frame
    tf_s71.word_wrap = True
    tf_s71.margin_left = Inches(0.3)
    tf_s71.margin_right = Inches(0.3)
    tf_s71.margin_top = Inches(0.3)

    ps71_0 = tf_s71.paragraphs[0]
    ps71_0.text = "📢 Advert Placement (/advertise)"
    ps71_0.font.size = Pt(20)
    ps71_0.font.bold = True
    ps71_0.font.color.rgb = COLOR_GOLD

    items_s71 = [
        ("Vertical 9:16 Format:", "Advertisers submit vertical images or video clips automatically cropped to 1080x1920."),
        ("Targeting Options:", "Choose display on Website, Mobile App, or Both."),
        ("Flexible Packages:", "Preset duration packages set by super admin (e.g. 7 Days, 14 Days, 30 Days)."),
        ("Display Frequency:", "Control ad appearance frequency (every 5m, 10m, 15m, 30m, or once daily).")
    ]
    for head, desc in items_s71:
        p = tf_s71.add_paragraph()
        p.space_before = Pt(12)
        r1 = p.add_run()
        r1.text = "• " + head + " "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_GOLD_SOFT
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_DIM

    card_s7_2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.8))
    card_s7_2.fill.solid()
    card_s7_2.fill.fore_color.rgb = COLOR_PANEL
    card_s7_2.line.color.rgb = COLOR_BORDER

    tf_s72 = card_s7_2.text_frame
    tf_s72.word_wrap = True
    tf_s72.margin_left = Inches(0.3)
    tf_s72.margin_right = Inches(0.3)
    tf_s72.margin_top = Inches(0.3)

    ps72_0 = tf_s72.paragraphs[0]
    ps72_0.text = "📊 Publisher Ad Manager (/ad-manager)"
    ps72_0.font.size = Pt(20)
    ps72_0.font.bold = True
    ps72_0.font.color.rgb = COLOR_GOLD

    items_s72 = [
        ("Tokenized Access Link:", "Approved advertisers receive a secure access link sent to their registered email."),
        ("Live Performance Stats:", "Track real-time ad impressions, clicks, and Click-Through-Rate (CTR)."),
        ("Daily Summary Emails:", "Automated daily email reports dispatch campaign analytics."),
        ("Create More Ads:", "Advertisers can launch additional ad campaigns directly inside their manager.")
    ]
    for head, desc in items_s72:
        p = tf_s72.add_paragraph()
        p.space_before = Pt(12)
        r1 = p.add_run()
        r1.text = "• " + head + " "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_GOLD_SOFT
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_DIM

    # -------------------------------------------------------------
    # SLIDE 8: Member Step-by-Step Guide
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "Member Guide", "How to Use the Website & Mobile App")

    step_w = Inches(3.64)
    step_h = Inches(4.8)

    steps_data = [
        ("Step 1: Access Platform", "🌐", [
            ("Visit Website:", "Open the website on any desktop or phone browser."),
            ("Download Mobile App:", "Get the app on Google Play Store for Android."),
            ("Dismissible Banner:", "Floating Google Play button available for instant app download.")
        ]),
        ("Step 2: Engage & Study", "📖", [
            ("Watch Media Reels:", "Swipe through video clips, worship moments, and testimonies."),
            ("Read / Listen to Bible:", "Open Bible tab, pick translation, and tap 🔊 Play Audio."),
            ("Join Live Broadcasts:", "Tap 'Watch Live' during service times.")
        ]),
        ("Step 3: Connect & Give", "🤝", [
            ("Prayer Requests:", "Submit prayer requests privately or on the public Prayer Wall."),
            ("Give Tithes & Offerings:", "Tap Give to pay online or upload bank receipt."),
            ("Place an Ad:", "Submit vertical ads to promote your business.")
        ])
    ]

    for i, (stitle, sicon, sitems) in enumerate(steps_data):
        lpos = Inches(0.8 + i * 3.95)
        box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, lpos, Inches(1.8), step_w, step_h)
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_PANEL
        box.line.color.rgb = COLOR_BORDER

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = sicon + " " + stitle
        p0.font.size = Pt(17)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_GOLD

        for head, desc in sitems:
            p = tf.add_paragraph()
            p.space_before = Pt(12)
            r1 = p.add_run()
            r1.text = "• " + head + " "
            r1.font.bold = True
            r1.font.size = Pt(12)
            r1.font.color.rgb = COLOR_GOLD_SOFT
            r2 = p.add_run()
            r2.text = desc
            r2.font.size = Pt(12)
            r2.font.color.rgb = COLOR_DIM

    # -------------------------------------------------------------
    # SLIDE 9: Pastoral & Administrative Tools
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "Administration & Growth", "Pastoral & Church Management Tools")

    card_s9_l = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    card_s9_l.fill.solid()
    card_s9_l.fill.fore_color.rgb = COLOR_PANEL
    card_s9_l.line.color.rgb = COLOR_BORDER

    tf_s9l = card_s9_l.text_frame
    tf_s9l.word_wrap = True
    tf_s9l.margin_left = Inches(0.3)
    tf_s9l.margin_right = Inches(0.3)
    tf_s9l.margin_top = Inches(0.3)

    ps9l0 = tf_s9l.paragraphs[0]
    ps9l0.text = "📈 Growth & Follow-Up Tracking"
    ps9l0.font.size = Pt(20)
    ps9l0.font.bold = True
    ps9l0.font.color.rgb = COLOR_GOLD

    items_s9l = [
        ("Service Attendance Logger:", "Track attendance split by gender (Males / Females) per service."),
        ("Newcomer Follow-Up Queue:", "Track first-time guests with status pipeline (New → Contacted → Returned)."),
        ("WhatsApp Tap-to-Chat:", "Instant WhatsApp links for pastoral follow-up directly from the dashboard."),
        ("Shareable CSV Reports:", "Generate Google-Forms style shareable CSV links for leadership reporting.")
    ]
    for head, desc in items_s9l:
        p = tf_s9l.add_paragraph()
        p.space_before = Pt(12)
        r1 = p.add_run()
        r1.text = "• " + head + " "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_GOLD_SOFT
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_DIM

    card_s9_r = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.8))
    card_s9_r.fill.solid()
    card_s9_r.fill.fore_color.rgb = COLOR_PANEL
    card_s9_r.line.color.rgb = COLOR_BORDER

    tf_s9r = card_s9_r.text_frame
    tf_s9r.word_wrap = True
    tf_s9r.margin_left = Inches(0.3)
    tf_s9r.margin_right = Inches(0.3)
    tf_s9r.margin_top = Inches(0.3)

    ps9r0 = tf_s9r.paragraphs[0]
    ps9r0.text = "🛡️ Security & Communications"
    ps9r0.font.size = Pt(20)
    ps9r0.font.bold = True
    ps9r0.font.color.rgb = COLOR_GOLD

    items_s9r = [
        ("Security Unblock PIN:", "4-6 digit PIN allows users/admins to self-unblock if account/IP is locked."),
        ("Email OTP Password Reset:", "6-digit OTP delivery to primary/backup emails with PIN fallback."),
        ("cPanel Corporate Email:", "Auto-creates corporate emails (e.g. admin@domain) on registration approval."),
        ("Push Notifications Wizard:", "Firebase Cloud Messaging integration for real-time app alerts.")
    ]
    for head, desc in items_s9r:
        p = tf_s9r.add_paragraph()
        p.space_before = Pt(12)
        r1 = p.add_run()
        r1.text = "• " + head + " "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_GOLD_SOFT
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_DIM

    # -------------------------------------------------------------
    # SLIDE 10: Conclusion & Call to Action
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)

    card10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1))
    card10.fill.solid()
    card10.fill.fore_color.rgb = COLOR_PANEL
    card10.line.color.rgb = COLOR_GOLD
    card10.line.width = Pt(1.5)

    tb10 = s10.shapes.add_textbox(Inches(1.8), Inches(1.6), Inches(9.733), Inches(4.2))
    tf10 = tb10.text_frame
    tf10.word_wrap = True

    p10_1 = tf10.paragraphs[0]
    p10_1.text = "GET STARTED TODAY"
    p10_1.font.size = Pt(14)
    p10_1.font.bold = True
    p10_1.font.color.rgb = COLOR_GOLD_SOFT
    p10_1.alignment = PP_ALIGN.CENTER

    p10_2 = tf10.add_paragraph()
    p10_2.text = "Download the Mobile App & Join Us Online!"
    p10_2.font.size = Pt(32)
    p10_2.font.bold = True
    p10_2.font.color.rgb = COLOR_WHITE
    p10_2.alignment = PP_ALIGN.CENTER
    p10_2.space_before = Pt(14)

    p10_3 = tf10.add_paragraph()
    p10_3.text = "1. Scan QR Code / Search on Google Play Store\n2. Explore Media Reels, Audio Bible & Events\n3. Share with Family, Friends & Church Members"
    p10_3.font.size = Pt(18)
    p10_3.font.color.rgb = COLOR_DIM
    p10_3.alignment = PP_ALIGN.CENTER
    p10_3.space_before = Pt(20)

    p10_4 = tf10.add_paragraph()
    p10_4.text = "Thank You & God Bless Our Ministry!"
    p10_4.font.size = Pt(22)
    p10_4.font.bold = True
    p10_4.font.color.rgb = COLOR_GOLD
    p10_4.alignment = PP_ALIGN.CENTER
    p10_4.space_before = Pt(24)

    prs.save("promo/Church_Media_Presentation.pptx")
    print("Successfully generated promo/Church_Media_Presentation.pptx")

if __name__ == "__main__":
    create_presentation()
